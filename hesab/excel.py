"""Write a Month to an .xlsx workbook: summary sheet + one RTL sheet per person."""

import io
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .build import Month, Person

MONTHS = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
          "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]

# (field, header, width)
COLUMNS = [
    ("date", "تاریخ", 10),
    ("weekday", "روز", 11),
    ("in1", "ورود۱", 9),
    ("out1", "خروج۱", 9),
    ("in2", "ورود۲", 9),
    ("out2", "خروج۲", 9),
    ("mission_from", "ماموریت از", 11),
    ("mission_to", "ماموریت تا", 11),
    ("remote", "ساعت دورکاری", 12),
    ("total", "جمع کارکرد روزانه", 16),
    ("notes", "توضیحات", 45),
]
COL = {f: get_column_letter(i + 1) for i, (f, _, _) in enumerate(COLUMNS)}

FONT = "Vazirmatn"
HEADER_ROW = 4
FIRST_ROW = HEADER_ROW + 1

thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")
FRIDAY_FILL = PatternFill("solid", fgColor="D9D9D9")
FLAG_FILL = PatternFill("solid", fgColor="FFF2A8")
MISSING_FILL = PatternFill("solid", fgColor="F4F4F4")
TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _sheet_title(name: str, used: set[str]) -> str:
    t = re.sub(r"[\[\]:*?/\\]", "", name).strip()[:31] or "sheet"
    base, i = t, 2
    while t in used:
        t = f"{base[:28]} {i}"
        i += 1
    used.add(t)
    return t


def _header(ws, title: str, subtitle: str, ncols: int) -> None:
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, size=12, bold=True)
    for c in ("A1", "A2"):
        ws[c].alignment = CENTER


def _pair(c1: str, c2: str, r: int) -> str:
    return f'IF(AND({c1}{r}<>"",{c2}{r}<>""),MOD({c2}{r}-{c1}{r},1),0)'


def _person_sheet(ws, p: Person, month: Month) -> int:
    """Fill a person's sheet; returns the totals row number."""
    ws.sheet_view.rightToLeft = True
    period = f"{MONTHS[month.month - 1]} {month.year or ''}".strip()
    _header(ws, "فرم حضور و غیاب کارکنان — تاف فناور پارس", f"{p.name} — {period}", len(COLUMNS))

    for i, (_, header, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(HEADER_ROW, i, header)
        c.font = Font(name=FONT, bold=True)
        c.fill, c.alignment, c.border = HEADER_FILL, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[HEADER_ROW].height = 30

    r = FIRST_ROW
    for d in p.days:
        values = {
            "date": f"{month.month:02d}/{d.day:02d}",
            "weekday": d.weekday,
            **{f: (v % 1440) / 1440 for f, v in d.times.items() if v is not None},
            "remote": d.remote / 1440 if d.remote is not None else None,
            "total": "=" + "+".join([_pair(COL["in1"], COL["out1"], r),
                                     _pair(COL["in2"], COL["out2"], r),
                                     _pair(COL["mission_from"], COL["mission_to"], r),
                                     f'N({COL["remote"]}{r})']),
            "notes": " | ".join(d.notes) or None,
        }
        friday = d.weekday == "جمعه"
        for i, (f, _, _) in enumerate(COLUMNS, start=1):
            c = ws.cell(r, i, values.get(f))
            c.font = Font(name=FONT, bold=(f == "total"))
            c.border = BORDER
            c.alignment = RIGHT if f == "notes" else CENTER
            if f in ("remote", "total"):
                c.number_format = "[h]:mm"
            elif f not in ("date", "weekday", "notes"):
                c.number_format = "h:mm"
            if f in d.flags:
                c.fill = FLAG_FILL
                c.comment = Comment("مقدار نامطمئن — با برگه اصلی چک شود", "hesab")
            elif friday:
                c.fill = FRIDAY_FILL
            elif not d.on_page:
                c.fill = MISSING_FILL
        r += 1

    last = r - 1
    total_row = r
    m_from, m_to = COL["mission_from"], COL["mission_to"]
    totals = {
        "date": "جمع ماه",
        "mission_to": (f'=SUMPRODUCT(({m_from}{FIRST_ROW}:{m_from}{last}<>"")*({m_to}{FIRST_ROW}:{m_to}{last}<>"")'
                       f'*MOD({m_to}{FIRST_ROW}:{m_to}{last}-{m_from}{FIRST_ROW}:{m_from}{last},1))'),
        "remote": f"=SUM({COL['remote']}{FIRST_ROW}:{COL['remote']}{last})",
        "total": f"=SUM({COL['total']}{FIRST_ROW}:{COL['total']}{last})",
    }
    for i, (f, _, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(total_row, i, totals.get(f))
        c.font = Font(name=FONT, bold=True)
        c.fill, c.border, c.alignment = TOTAL_FILL, BORDER, CENTER
        if f in ("mission_to", "remote", "total"):
            c.number_format = "[h]:mm"
    ws.merge_cells(f"A{total_row}:B{total_row}")
    ws.cell(total_row, 7, "جمع ماموریت").font = Font(name=FONT, bold=True)
    ws.freeze_panes = ws.cell(FIRST_ROW, 3)
    ws.print_options.horizontalCentered = True
    ws.page_setup.fitToWidth = 1
    return total_row


SUMMARY_COLUMNS = [
    ("نام و نام خانوادگی", 26),
    ("روزهای حضور", 13),
    ("جمع کارکرد", 13),
    ("جمع ماموریت", 13),
    ("جمع دورکاری", 13),
    ("کارکرد جمعه", 13),
    ("موارد نیازمند بررسی", 18),
]


def write_workbook(month: Month) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "خلاصه"
    summary.sheet_view.rightToLeft = True
    period = f"{MONTHS[month.month - 1]} {month.year or ''}".strip()
    _header(summary, "خلاصه کارکرد ماهانه — تاف فناور پارس", period, len(SUMMARY_COLUMNS))
    for i, (h, w) in enumerate(SUMMARY_COLUMNS, start=1):
        c = summary.cell(HEADER_ROW, i, h)
        c.font = Font(name=FONT, bold=True)
        c.fill, c.alignment, c.border = HEADER_FILL, CENTER, BORDER
        summary.column_dimensions[get_column_letter(i)].width = w

    used = {"خلاصه"}
    for n, p in enumerate(month.people):
        ws = wb.create_sheet(_sheet_title(p.name, used))
        tr = _person_sheet(ws, p, month)
        ref = f"'{ws.title}'!"
        rng = lambda col: f"{ref}{col}{FIRST_ROW}:{col}{tr - 1}"  # noqa: E731
        flagged = sum(1 for d in p.days if d.flags)
        row = [
            f'=HYPERLINK("#\'{ws.title}\'!A1","{p.name}")',
            f'=COUNTIF({rng(COL["total"])},">0")',
            f"={ref}{COL['total']}{tr}",
            f"={ref}{COL['mission_to']}{tr}",
            f"={ref}{COL['remote']}{tr}",
            f'=SUMIF({rng(COL["weekday"])},"جمعه",{rng(COL["total"])})',
            flagged,
        ]
        r = FIRST_ROW + n
        for i, v in enumerate(row, start=1):
            c = summary.cell(r, i, v)
            c.font = Font(name=FONT, bold=(i == 1), color="1F4E79" if i == 1 else None)
            c.border, c.alignment = BORDER, CENTER
            if 3 <= i <= 6:
                c.number_format = "[h]:mm"
            if i == 7 and flagged:
                c.fill = FLAG_FILL

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
